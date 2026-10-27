import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill
from datetime import date, datetime
import os, re

st.set_page_config(page_title="SUNMIT Travel | Made Simple", page_icon="🚌", layout="wide")

FILE_NAME = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sunmit_travel_bookings.xlsx")

CITIES = ["Pune", "Mumbai", "Nashik", "Shirpur", "Dhule", "Indore", "Surat", "Ahmedabad", "Aurangabad", "Nagpur"]
BUSES = {
    "SUNMIT Express - AC Seater": {"type": "AC Seater", "fare": 650},
    "Royal Roadlines - Non-AC Seater": {"type": "Non-AC Seater", "fare": 350},
    "BlueLine Travels - AC Sleeper": {"type": "AC Sleeper", "fare": 950},
    "Maharaja Coaches - Volvo Sleeper": {"type": "Volvo AC Sleeper", "fare": 1350},
    "Night Rider - Non-AC Sleeper": {"type": "Non-AC Sleeper", "fare": 600},
}
HEADERS = ["Booking ID", "Booked At", "Journey Date", "From City", "To City", "Bus Name", "Bus Type", "Passenger Name", "Contact Number", "Address", "Seat Numbers", "Passenger Genders", "Couple Booking", "Passengers", "Fare Per Seat", "Total Amount", "Status"]

def setup_excel():
    if not os.path.exists(FILE_NAME):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Bookings"
        ws.append(HEADERS)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="16365C")
        wb.save(FILE_NAME)
    else:
        wb = openpyxl.load_workbook(FILE_NAME)
        if "Bookings" not in wb.sheetnames:
            ws = wb.create_sheet("Bookings", 0)
            ws.append(HEADERS)
            wb.save(FILE_NAME)

setup_excel()

def load_data():
    try:
        df = pd.read_excel(FILE_NAME, sheet_name="Bookings")
        return df
    except Exception:
        return pd.DataFrame(columns=HEADERS)

def save_new_booking(row_data):
    wb = openpyxl.load_workbook(FILE_NAME)
    ws = wb["Bookings"]
    ws.append(row_data)
    wb.save(FILE_NAME)

def cancel_booking_id(bid):
    wb = openpyxl.load_workbook(FILE_NAME)
    ws = wb["Bookings"]
    found = False
    for i in range(2, ws.max_row + 1):
        if str(ws.cell(i, 1).value) == bid:
            ws.cell(i, 17).value = "Cancelled"
            found = True
            break
    if found:
        wb.save(FILE_NAME)
    return found

df = load_data()

# Sidebar Navigation
st.sidebar.title("🚌 SUNMIT Travel")
st.sidebar.caption("M A D E   S I M P L E")
menu = st.sidebar.radio("Navigation", ["▣ Book Ticket", "⌂ Dashboard", "▤ Bookings", "▥ Reports", "⚙ Settings"])

# Metrics / Top stats
active_df = df[df["Status"].astype(str) == "Confirmed"] if not df.empty and "Status" in df.columns else pd.DataFrame()
today_str = date.today().isoformat()
today_bookings = active_df[active_df["Booked At"].astype(str).str.startswith(today_str)] if not active_df.empty else pd.DataFrame()
total_revenue = active_df["Total Amount"].sum() if not active_df.empty else 0

col1, col2, col3 = st.columns(3)
col1.metric("CONFIRMED BOOKINGS", len(active_df))
col2.metric("BOOKINGS TODAY", len(today_bookings))
col3.metric("ACTIVE REVENUE", f"₹{total_revenue:,.0f}")
st.divider()

# Page 1: Book Ticket
if menu == "▣ Book Ticket":
    st.subheader("Book a Ticket")
    c1, c2 = st.columns([1, 1])

    with c1:
        st.markdown("### Passenger Details")
        p_name = st.text_input("Passenger Name *")
        p_contact = st.text_input("Mobile Number * (10-13 digits)")
        p_address = st.text_input("Address")
        p_gender = st.selectbox("Passenger Gender", ["Female", "Male", "Prefer not to say"])
        p_couple = st.checkbox("Couple booking (husband/wife or partners)")

        st.markdown("### Journey Details")
        src = st.selectbox("From City", CITIES, index=0)
        dst_options = [c for c in CITIES if c != src]
        dst = st.selectbox("To City", dst_options, index=0)
        jdate = st.date_input("Journey Date", min_value=date.today())
        bus_name = st.selectbox("Bus Operator", list(BUSES.keys()))
        bus_info = BUSES[bus_name]
        p_count = st.number_input("Passenger Count", min_value=1, max_value=6, value=1)

    with c2:
        st.markdown("### Seat Selection")
        # Occupied seats filter
        occupied_seats = set()
        if not df.empty and "Status" in df.columns:
            matches = df[(df["Journey Date"].astype(str) == str(jdate)) & 
                         (df["From City"] == src) & 
                         (df["To City"] == dst) & 
                         (df["Bus Name"] == bus_name) & 
                         (df["Status"] == "Confirmed")]
            for seat_str in matches["Seat Numbers"].dropna():
                for s in str(seat_str).split(","):
                    if s.strip():
                        occupied_seats.add(s.strip())

        # Generate seat lists
        if "Sleeper" in bus_info["type"]:
            all_seats = [f"L{r}" for r in range(1, 9)] + [f"U{r}" for r in range(1, 9)] + [f"S{r}" for r in range(1, 9)]
        else:
            all_seats = [f"{ch}{r}" for r in range(1, 11) for ch in ["A", "B", "C", "D"]]

        available_seats = [s for s in all_seats if s not in occupied_seats]
        if occupied_seats:
            st.info(f"Already Booked: {', '.join(sorted(occupied_seats))}")

        selected_seats = st.multiselect("Choose Seats:", available_seats, max_selections=int(p_count))

        st.markdown("### Booking Summary")
        fare_per_seat = bus_info["fare"]
        total_fare = fare_per_seat * int(p_count)

        st.write(f"**Route:** {src} → {dst}")
        st.write(f"**Bus Type:** {bus_info['type']}")
        st.write(f"**Selected Seats:** {', '.join(selected_seats) if selected_seats else 'None'}")
        st.write(f"**Fare Per Seat:** ₹{fare_per_seat}")
        st.success(f"### Total Amount: ₹{total_fare:,.0f}")

        if st.button("Confirm Booking", type="primary", use_container_width=True):
            if not p_name.strip() or not re.match(r"^[a-zA-Z\s\.\-\']+$", p_name):
                st.error("Name must contain letters only.")
            elif not p_contact.isdigit() or not (10 <= len(p_contact) <= 13):
                st.error("Enter a valid mobile number (10 to 13 digits).")
            elif len(selected_seats) != p_count:
                st.error(f"Please select exactly {p_count} seat(s).")
            elif p_couple and p_count < 2:
                st.error("Select at least 2 passengers for couple booking.")
            else:
                bid = "SUN" + datetime.now().strftime("%y%m%d%H%M%S")
                booked_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                genders = [p_gender if not p_couple or i % 2 == 0 else "Partner" for i in range(p_count)]
                
                row = [
                    bid, booked_time, str(jdate), src, dst, bus_name,
                    bus_info["type"], p_name.strip(), p_contact.strip(),
                    p_address.strip(), ", ".join(selected_seats),
                    ", ".join(genders), "Yes" if p_couple else "No",
                    int(p_count), fare_per_seat, total_fare, "Confirmed"
                ]
                save_new_booking(row)
                st.balloons()
                st.success(f"Booking Confirmed! ID: {bid}")
                st.rerun()

# Page 2: Dashboard
elif menu == "⌂ Dashboard":
    st.subheader("Recent Bookings")
    if not df.empty:
        disp_cols = ["Booking ID", "Passenger Name", "From City", "To City", "Journey Date", "Bus Name", "Total Amount", "Status"]
        existing = [c for c in disp_cols if c in df.columns]
        st.dataframe(df[existing].tail(15).iloc[::-1], use_container_width=True)
    else:
        st.info("No bookings recorded yet.")

# Page 3: Bookings Management
elif menu == "▤ Bookings":
    st.subheader("Booking Management")
    if not df.empty:
        search_txt = st.text_input("Search bookings by ID, Name or Mobile")
        filtered_df = df
        if search_txt:
            filtered_df = df[df.astype(str).apply(lambda row: search_txt.lower() in row.to_string().lower(), axis=1)]
        
        st.dataframe(filtered_df, use_container_width=True)

        st.divider()
        st.markdown("### Cancel a Booking")
        col_c1, col_c2 = st.columns([2, 1])
        with col_c1:
            cancel_id = st.text_input("Enter Booking ID to Cancel (e.g. SUN...)")
        with col_c2:
            st.write("")
            st.write("")
            if st.button("Cancel Booking", type="secondary", use_container_width=True):
                if cancel_id:
                    if cancel_booking_id(cancel_id.strip()):
                        st.success(f"Booking {cancel_id} has been cancelled.")
                        st.rerun()
                    else:
                        st.error("Booking ID not found.")
                else:
                    st.warning("Please enter a Booking ID.")
    else:
        st.info("No bookings found.")

# Page 4: Reports
elif menu == "▥ Reports":
    st.subheader("Revenue by Bus Operator")
    if not active_df.empty:
        rev_by_bus = active_df.groupby("Bus Name")["Total Amount"].sum().reset_index()
        rev_by_bus.columns = ["Bus Operator", "Revenue (₹)"]
        st.dataframe(rev_by_bus, use_container_width=True)
        st.bar_chart(rev_by_bus.set_index("Bus Operator"))
    else:
        st.info("No confirmed booking data available for reports.")

# Page 5: Settings
elif menu == "⚙ Settings":
    st.subheader("System Information")
    st.markdown(f"""
    - **System:** SUNMIT Travel — Made Simple (Streamlit Cloud Edition)
    - **Data Store:** `{FILE_NAME}`
    - **Status:** Cloud Hosted & Connected to Excel
    """)
